from io import BytesIO
import os
from pathlib import Path
from typing import Any
from fastapi import Depends, FastAPI, File, HTTPException, Response, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel
from sqlalchemy.orm import Session
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from .database import Base, SessionLocal, engine, get_db
from .models import Aluno, Resultado, Turma, Usuario
from .scanner import ScanError, scan_card
from .auth import admin_user, can_access_turma, clear_login_cookie, current_user, hash_password, set_login_cookie, verify_password

BASE_DIR = Path(__file__).resolve().parent
Base.metadata.create_all(bind=engine)

def bootstrap_admin():
    login = os.getenv('ADMIN_LOGIN', '').strip().lower()
    password = os.getenv('ADMIN_PASSWORD', '')
    name = os.getenv('ADMIN_NAME', 'Administrador').strip() or 'Administrador'
    if not login or not password: return
    db = SessionLocal()
    try:
        if not db.query(Usuario).filter(Usuario.login == login).first():
            db.add(Usuario(nome=name, login=login, senha_hash=hash_password(password), tipo='admin', ativo=True)); db.commit()
    finally: db.close()
bootstrap_admin()

app = FastAPI(title='EDUSCANNER')
app.mount('/static', StaticFiles(directory=BASE_DIR / 'static'), name='static')
@app.get('/')
def home(): return FileResponse(BASE_DIR / 'static' / 'index.html')
@app.get('/health')
def health(): return {'ok': True}

class LoginIn(BaseModel): login: str; senha: str
class UserIn(BaseModel): nome: str; login: str; senha: str; tipo: str = 'professor'; turma_ids: list[int] = []
class UserUpdate(BaseModel): nome: str; login: str; senha: str | None = None; tipo: str = 'professor'; ativo: bool = True; turma_ids: list[int] = []

@app.post('/api/auth/login')
def login(payload: LoginIn, response: Response, db: Session = Depends(get_db)):
    user = db.query(Usuario).filter(Usuario.login == payload.login.strip().lower()).first()
    if not user or not user.ativo or not verify_password(payload.senha, user.senha_hash): raise HTTPException(401, 'Usuário ou senha inválidos.')
    set_login_cookie(response, user.id); return {'ok': True, 'nome': user.nome, 'tipo': user.tipo}
@app.post('/api/auth/logout')
def logout(response: Response): clear_login_cookie(response); return {'ok': True}
@app.get('/api/auth/me')
def me(user: Usuario = Depends(current_user)): return {'id': user.id, 'nome': user.nome, 'login': user.login, 'tipo': user.tipo}

@app.post('/api/scan')
async def scan(file: UploadFile = File(...), user: Usuario = Depends(current_user)):
    try: return scan_card(await file.read())
    except ScanError as exc: return JSONResponse(status_code=422, content={'detail': str(exc)})
    except Exception: return JSONResponse(status_code=500, content={'detail': 'Não foi possível processar a foto. Tente outra imagem.'})

@app.get('/api/turmas')
def listar_turmas(db: Session = Depends(get_db), user: Usuario = Depends(current_user)):
    turmas = db.query(Turma).order_by(Turma.nome).all() if user.tipo == 'admin' else sorted(user.turmas, key=lambda t:t.nome)
    return [{'id': t.id, 'nome': t.nome, 'total_alunos': len(t.alunos)} for t in turmas]
@app.get('/api/turmas/{turma_id}/alunos')
def listar_alunos(turma_id: int, db: Session = Depends(get_db), user: Usuario = Depends(current_user)):
    turma = db.get(Turma, turma_id)
    if not turma: raise HTTPException(404, 'Turma não encontrada.')
    if not can_access_turma(user, turma_id): raise HTTPException(403, 'Você não tem acesso a esta turma.')
    alunos = db.query(Aluno).filter(Aluno.turma_id == turma_id).order_by(Aluno.nome).all()
    return [{'id': a.id, 'nome': a.nome, 'matricula': a.matricula} for a in alunos]

@app.post('/api/alunos/importar')
async def importar_alunos(file: UploadFile = File(...), db: Session = Depends(get_db), user: Usuario = Depends(admin_user)):
    if not (file.filename or '').lower().endswith(('.xlsx', '.xlsm')): raise HTTPException(400, 'Envie uma planilha .xlsx.')
    try:
        wb=load_workbook(BytesIO(await file.read()),read_only=True,data_only=True); rows=list(wb.active.iter_rows(values_only=True))
    except Exception: raise HTTPException(400, 'Não foi possível ler a planilha.')
    if not rows: raise HTTPException(400, 'A planilha está vazia.')
    def norm(v): return str(v or '').strip().lower().replace('í','i').replace('á','a').replace('ã','a').replace('ç','c')
    headers=[norm(v) for v in rows[0]]; aliases={'nome':['nome','nome do aluno','aluno'],'turma':['turma','sala'],'matricula':['matricula','matrícula','registro','ra']}; indexes={}
    for field,names in aliases.items():
        for i,h in enumerate(headers):
            if h in [norm(n) for n in names]: indexes[field]=i; break
    if len(indexes)!=3: raise HTTPException(400,'A primeira linha precisa conter as colunas Nome do aluno, Turma e Matrícula.')
    criados=atualizados=ignorados=0
    for row in rows[1:]:
        def value(field):
            i=indexes[field]; return str(row[i] if i<len(row) and row[i] is not None else '').strip()
        nome,turma_nome,matricula=value('nome'),value('turma'),value('matricula')
        if not nome or not turma_nome or not matricula: ignorados+=1; continue
        turma=db.query(Turma).filter(Turma.nome==turma_nome).first()
        if not turma: turma=Turma(nome=turma_nome); db.add(turma); db.flush()
        aluno=db.query(Aluno).filter(Aluno.matricula==matricula).first()
        if aluno: aluno.nome=nome; aluno.turma_id=turma.id; atualizados+=1
        else: db.add(Aluno(nome=nome,matricula=matricula,turma_id=turma.id)); criados+=1
    db.commit(); return {'criados':criados,'atualizados':atualizados,'ignorados':ignorados}

@app.get('/api/admin/usuarios')
def list_users(db: Session=Depends(get_db), _:Usuario=Depends(admin_user)):
    users=db.query(Usuario).order_by(Usuario.nome).all(); return [{'id':u.id,'nome':u.nome,'login':u.login,'tipo':u.tipo,'ativo':u.ativo,'turma_ids':[t.id for t in u.turmas]} for u in users]
@app.post('/api/admin/usuarios')
def create_user(payload:UserIn, db:Session=Depends(get_db), _:Usuario=Depends(admin_user)):
    login=payload.login.strip().lower()
    if len(payload.senha)<6: raise HTTPException(400,'A senha deve ter pelo menos 6 caracteres.')
    if payload.tipo not in ('admin','professor'): raise HTTPException(400,'Tipo inválido.')
    if db.query(Usuario).filter(Usuario.login==login).first(): raise HTTPException(409,'Este login já está cadastrado.')
    turmas=db.query(Turma).filter(Turma.id.in_(payload.turma_ids)).all() if payload.turma_ids else []
    u=Usuario(nome=payload.nome.strip(),login=login,senha_hash=hash_password(payload.senha),tipo=payload.tipo,ativo=True,turmas=turmas); db.add(u); db.commit(); db.refresh(u); return {'id':u.id}
@app.put('/api/admin/usuarios/{user_id}')
def update_user(user_id:int,payload:UserUpdate,db:Session=Depends(get_db),admin:Usuario=Depends(admin_user)):
    u=db.get(Usuario,user_id)
    if not u: raise HTTPException(404,'Usuário não encontrado.')
    login=payload.login.strip().lower(); existing=db.query(Usuario).filter(Usuario.login==login,Usuario.id!=user_id).first()
    if existing: raise HTTPException(409,'Este login já está cadastrado.')
    if payload.tipo not in ('admin','professor'): raise HTTPException(400,'Tipo inválido.')
    if user_id==admin.id and (payload.tipo!='admin' or not payload.ativo): raise HTTPException(400,'Você não pode remover seu próprio acesso de administrador.')
    u.nome=payload.nome.strip();u.login=login;u.tipo=payload.tipo;u.ativo=payload.ativo
    if payload.senha:
        if len(payload.senha)<6: raise HTTPException(400,'A senha deve ter pelo menos 6 caracteres.')
        u.senha_hash=hash_password(payload.senha)
    u.turmas=db.query(Turma).filter(Turma.id.in_(payload.turma_ids)).all() if payload.turma_ids else []
    db.commit(); return {'ok':True}



def _pdf_resultado(r: Resultado, turma: Turma) -> BytesIO:
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, rightMargin=16*mm, leftMargin=16*mm, topMargin=14*mm, bottomMargin=14*mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle('SchoolTitle', parent=styles['Title'], alignment=TA_CENTER, fontSize=15, leading=18, spaceAfter=3)
    sub = ParagraphStyle('SchoolSub', parent=styles['Normal'], alignment=TA_CENTER, fontSize=9, textColor=colors.HexColor('#555555'))
    story = []
    logo = BASE_DIR / 'static' / 'logo-escola.png'
    if logo.exists():
        story += [RLImage(str(logo), width=35*mm, height=35*mm), Spacer(1, 2*mm)]
    story += [Paragraph('Colégio Estadual em Período Integral João Barbosa Reis', title), Paragraph('CEPI-JBR - Relatório de Resultado', sub), Spacer(1, 7*mm)]
    info = [
        ['Aluno', r.aluno.nome], ['Matrícula', r.aluno.matricula], ['Turma', turma.nome],
        ['Prova / bloco', r.prova_nome], ['Questões', str(r.quantidade_questoes)],
        ['Acertos', str(r.acertos)], ['Erros', str(r.erros)], ['Anuladas', str(r.anuladas)],
        ['Em branco', str(r.em_branco)], ['Nota', f'{r.nota:.2f}'.replace('.', ',')],
        ['Data', r.criado_em.strftime('%d/%m/%Y %H:%M')],
    ]
    table = Table(info, colWidths=[42*mm, 120*mm])
    table.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(0,-1),colors.HexColor('#f1f3f5')), ('FONTNAME',(0,0),(0,-1),'Helvetica-Bold'),
        ('GRID',(0,0),(-1,-1),0.5,colors.HexColor('#d6d9dc')), ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('LEFTPADDING',(0,0),(-1,-1),7), ('RIGHTPADDING',(0,0),(-1,-1),7), ('TOPPADDING',(0,0),(-1,-1),7), ('BOTTOMPADDING',(0,0),(-1,-1),7),
    ]))
    story += [table, Spacer(1, 8*mm), Paragraph('Documento gerado pelo EduScanner. Relatório para conferência e uso administrativo.', sub)]
    doc.build(story)
    out.seek(0)
    return out

@app.delete('/api/admin/usuarios/{user_id}')
def delete_user(user_id:int, db:Session=Depends(get_db), admin:Usuario=Depends(admin_user)):
    u=db.get(Usuario,user_id)
    if not u: raise HTTPException(404,'Usuário não encontrado.')
    if u.id==admin.id: raise HTTPException(400,'Você não pode excluir o usuário administrador que está em uso.')
    if u.tipo!='professor': raise HTTPException(400,'Por segurança, somente cadastros de professor podem ser excluídos por esta opção.')
    db.delete(u); db.commit(); return {'ok':True}

@app.delete('/api/resultados/{resultado_id}')
def delete_resultado(resultado_id:int, db:Session=Depends(get_db), user:Usuario=Depends(current_user)):
    r=db.get(Resultado,resultado_id)
    if not r: raise HTTPException(404,'Resultado não encontrado.')
    if not can_access_turma(user,r.aluno.turma_id): raise HTTPException(403,'Você não tem acesso a este resultado.')
    db.delete(r); db.commit(); return {'ok':True}

@app.get('/api/resultados/{resultado_id}/pdf')
def resultado_pdf(resultado_id:int, db:Session=Depends(get_db), user:Usuario=Depends(current_user)):
    r=db.get(Resultado,resultado_id)
    if not r: raise HTTPException(404,'Resultado não encontrado.')
    if not can_access_turma(user,r.aluno.turma_id): raise HTTPException(403,'Você não tem acesso a este resultado.')
    pdf=_pdf_resultado(r,r.aluno.turma)
    safe=''.join(c if c.isalnum() or c in '-_' else '_' for c in f'{r.aluno.nome}_{r.prova_nome}')[:100]
    return StreamingResponse(pdf, media_type='application/pdf', headers={'Content-Disposition':f'attachment; filename="relatorio_{safe}.pdf"'})

class ResultadoIn(BaseModel):
    aluno_id:int; prova_nome:str='Prova'; quantidade_questoes:int; gabarito:list[str]; questions:list[dict[str,Any]]
@app.post('/api/resultados')
def salvar_resultado(payload:ResultadoIn,db:Session=Depends(get_db),user:Usuario=Depends(current_user)):
    aluno=db.get(Aluno,payload.aluno_id)
    if not aluno: raise HTTPException(404,'Aluno não encontrado.')
    if not can_access_turma(user,aluno.turma_id): raise HTTPException(403,'Você não tem acesso a este aluno.')
    count=payload.quantidade_questoes
    if count not in (20,30,45) or len(payload.gabarito)!=count: raise HTTPException(400,'Quantidade/gabarito inválido.')
    qs=sorted(payload.questions,key=lambda q:int(q.get('number',0)))[:count]
    if len(qs)<count: raise HTTPException(400,'Leitura incompleta do cartão.')
    correct=mult=blank=0;answers=[]
    for i,q in enumerate(qs):
        status=q.get('status');answer=q.get('answer') if status=='OK' else None;answers.append(answer or ('MULT' if status=='MULT' else ''))
        if status=='MULT':mult+=1
        elif status=='BLANK':blank+=1
        if status=='OK' and answer==payload.gabarito[i]:correct+=1
    wrong=count-correct;nota=round(correct/count*10,2)
    result=Resultado(aluno_id=aluno.id,prova_nome=payload.prova_nome.strip() or 'Prova',quantidade_questoes=count,gabarito=','.join(payload.gabarito),respostas=','.join(answers),acertos=correct,erros=wrong,anuladas=mult,em_branco=blank,nota=nota)
    db.add(result);db.commit();db.refresh(result);return {'id':result.id,'acertos':correct,'erros':wrong,'anuladas':mult,'em_branco':blank,'nota':nota}
@app.get('/api/turmas/{turma_id}/relatorio')
def relatorio(turma_id:int,prova_nome:str|None=None,db:Session=Depends(get_db),user:Usuario=Depends(current_user)):
    turma=db.get(Turma,turma_id)
    if not turma:raise HTTPException(404,'Turma não encontrada.')
    if not can_access_turma(user,turma_id):raise HTTPException(403,'Você não tem acesso a esta turma.')
    query=db.query(Resultado).join(Aluno).filter(Aluno.turma_id==turma_id)
    if prova_nome:query=query.filter(Resultado.prova_nome==prova_nome)
    resultados=query.order_by(Aluno.nome,Resultado.criado_em.desc()).all()
    return {'turma':turma.nome,'resultados':[{'id':r.id,'aluno':r.aluno.nome,'matricula':r.aluno.matricula,'prova':r.prova_nome,'questoes':r.quantidade_questoes,'acertos':r.acertos,'erros':r.erros,'anuladas':r.anuladas,'em_branco':r.em_branco,'nota':r.nota,'data':r.criado_em.isoformat()} for r in resultados]}
